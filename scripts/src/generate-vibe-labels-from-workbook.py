#!/usr/bin/env python3
"""Generate the version-controlled VIBES taxonomy from an approved workbook.

The workbook is an input artifact and is deliberately not copied into the source
repository. Review its Vibe Reference tab before running this generator; only
explicit rows with a category, label, and helper text are emitted. Rows whose
label begins `N/A` preserve the workbook's endorsement-only rule by not making
that category VIBE-eligible.
"""
from __future__ import annotations

import argparse
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook


def quote(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("lib/constants/src/vibe-labels.ts"),
    )
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    if "Vibe Reference" not in workbook.sheetnames:
        raise ValueError("Approved workbook does not contain a Vibe Reference sheet")

    taxonomy: OrderedDict[str, list[tuple[str, str]]] = OrderedDict()
    for row_number, row in enumerate(workbook["Vibe Reference"].iter_rows(min_row=2, values_only=True), start=2):
        category, label, helper = row[:3]
        if not any((category, label, helper)):
            continue
        if not all((category, label, helper)):
            raise ValueError(f"Incomplete Vibe Reference row {row_number}")
        category_text = str(category).strip()
        label_text = str(label).strip()
        helper_text = str(helper).strip()
        if label_text.startswith("N/A"):
            continue
        values = taxonomy.setdefault(category_text, [])
        if label_text not in {entry[0] for entry in values}:
            values.append((label_text, helper_text))

    if not taxonomy:
        raise ValueError("No VIBE records were found in the approved workbook")

    lines = [
        "/**",
        " * Mapping With Melanin™ — Master VIBES taxonomy.",
        " *",
        " * Generated from the approved `Vibe Reference` worksheet. A VIBE",
        " * describes the atmosphere or visit experience, not a category, an",
        " * ownership designation, or a quality claim. Endorsement-only categories",
        " * are intentionally absent from this map.",
        " */",
        "",
        "export interface VibeLabel {",
        "  label: string;",
        "  helperText: string;",
        "}",
        "",
        "export const VIBES_BY_CATEGORY: Record<string, VibeLabel[]> = {",
    ]
    for category, vibes in taxonomy.items():
        lines.append(f"  {quote(category)}: [")
        for label, helper in vibes:
            lines.append(f"    {{ label: {quote(label)}, helperText: {quote(helper)} }},")
        lines.append("  ],")
    lines.extend([
        "};",
        "",
        "/** Flat list of every explicit VIBE label, for cross-client validation. */",
        "export const ALL_VIBE_LABELS: string[] = [",
        "  ...new Set(Object.values(VIBES_BY_CATEGORY).flatMap((vibes) => vibes.map((vibe) => vibe.label))),",
        "];",
        "",
        "/** Categories with approved VIBES; all other categories use endorsements only. */",
        "export const VIBE_ELIGIBLE_CATEGORIES: string[] = Object.keys(VIBES_BY_CATEGORY);",
        "",
        "export function isVibeEligible(category: string): boolean {",
        "  return VIBE_ELIGIBLE_CATEGORIES.includes(category);",
        "}",
        "",
    ])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "categories": len(taxonomy), "vibes": sum(len(vibes) for vibes in taxonomy.values())}))


if __name__ == "__main__":
    main()
