#!/usr/bin/env python3
"""Generate the versioned founder taxonomy contract from a VIBESNOW workbook.

The workbook is the governed content source. The generated TypeScript module is
committed so every client can consume the same immutable category, VIBE, and
quick-review vocabulary without loading a spreadsheet at runtime.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

SUBCATEGORY_SHEET: Final = "_Subcategory Lists"
VIBE_SHEET: Final = "_Vibe Lists"
QUICK_REVIEW_SHEET: Final = "_Quick Review Matrix"
ENDORSEMENTS_ONLY_MARKER: Final = "N/A — Use Endorsements, Not Vibes"
CATEGORY_IDS: Final = {
    "Food & Drink": "food-and-drink",
    "Beauty & Personal Care": "beauty-and-personal-care",
    "Health & Wellness": "health-and-wellness",
    "Shopping & Retail": "shopping-and-retail",
    "Travel & Hospitality": "travel-and-hospitality",
    "Arts, Culture & Entertainment": "arts-culture-and-entertainment",
    "Professional Services": "professional-services",
    "Home & Property Services": "home-and-property-services",
    "Automotive & Transportation": "automotive-and-transportation",
    "Events & Celebrations": "events-and-celebrations",
    "Education & Learning": "education-and-learning",
    "Children & Family": "children-and-family",
    "Community & Nonprofit": "community-and-nonprofit",
    "Faith & Spirituality": "faith-and-spirituality",
    "Media & Creative Services": "media-and-creative-services",
    "Sports & Recreation": "sports-and-recreation",
    "Pets & Animal Services": "pets-and-animal-services",
    "Technology & Digital Services": "technology-and-digital-services",
    "Financial & Business Services": "financial-and-business-services",
    "Legal & Government Services": "legal-and-government-services",
    "Agriculture & Specialty Producers": "agriculture-and-specialty-producers",
    "Other Services": "other-services",
}


def read_columns(workbook_path: Path, sheet_name: str) -> dict[str, list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers: list[tuple[int, str]] = [
            (column, str(cell.value).strip())
            for column, cell in enumerate(sheet[1], start=1)
            if cell.value is not None and str(cell.value).strip()
        ]
        if not headers:
            raise ValueError(f"{sheet_name} has no category headers")

        columns: dict[str, list[str]] = {}
        for column, header in headers:
            values = [
                str(sheet.cell(row, column).value).strip()
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row, column).value is not None
                and str(sheet.cell(row, column).value).strip()
            ]
            columns[header] = values
        return columns
    finally:
        workbook.close()


def build_contract(workbook_path: Path) -> list[dict[str, object]]:
    subcategories = read_columns(workbook_path, SUBCATEGORY_SHEET)
    vibes = read_columns(workbook_path, VIBE_SHEET)
    quick_reviews = read_columns(workbook_path, QUICK_REVIEW_SHEET)

    category_order = list(subcategories)
    expected_categories = set(category_order)
    for name, values in ((VIBE_SHEET, vibes), (QUICK_REVIEW_SHEET, quick_reviews)):
        if set(values) != expected_categories:
            missing = sorted(expected_categories - set(values))
            unexpected = sorted(set(values) - expected_categories)
            raise ValueError(
                f"{name} category headings do not match {SUBCATEGORY_SHEET}; "
                f"missing={missing}, unexpected={unexpected}"
            )

    contract: list[dict[str, object]] = []
    for category in category_order:
        category_id = CATEGORY_IDS.get(category)
        if not category_id:
            raise ValueError(
                f"No stable category identifier registered for {category!r}. "
                "Founder approval is required before adding a new category."
            )
        source_vibes = vibes[category]
        if source_vibes == [ENDORSEMENTS_ONLY_MARKER]:
            vibe_mode = "endorsements_only"
            vibe_labels: list[str] = []
        elif ENDORSEMENTS_ONLY_MARKER in source_vibes:
            raise ValueError(
                f"{category} mixes the endorsements-only marker with VIBE labels"
            )
        else:
            vibe_mode = "vibes"
            vibe_labels = source_vibes

        contract.append(
            {
                "id": category_id,
                "category": category,
                "subcategories": subcategories[category],
                "vibeMode": vibe_mode,
                "vibeLabels": vibe_labels,
                "quickReviewLabels": quick_reviews[category],
            }
        )
    return contract


def render_typescript(workbook_path: Path, contract: list[dict[str, object]]) -> str:
    checksum = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    payload = json.dumps(contract, ensure_ascii=False, indent=2)
    return f'''/**
 * GENERATED FILE — DO NOT EDIT BY HAND.
 *
 * Generated from {workbook_path.name} with
 * scripts/generate-founder-taxonomy-contract.py.
 *
 * This is the versioned, founder-controlled source for business category /
 * subcategory relationships, category-permitted VIBES, and quick-review
 * endorsement wording. It contains no database identifiers or member content.
 */

export const FOUNDER_TAXONOMY_VERSION = "vibesnow-2026-10-15";
export const FOUNDER_TAXONOMY_WORKBOOK_SHA256 = "{checksum}";
export const FOUNDER_TAXONOMY_ENDORSEMENTS_ONLY_MARKER = "{ENDORSEMENTS_ONLY_MARKER}";

export type FounderVibeMode = "vibes" | "endorsements_only";

export interface FounderTaxonomyCategory {{
  readonly id: string;
  readonly category: string;
  readonly subcategories: readonly string[];
  readonly vibeMode: FounderVibeMode;
  readonly vibeLabels: readonly string[];
  readonly quickReviewLabels: readonly string[];
}}

export const FOUNDER_TAXONOMY_CONTRACT = {payload} as const satisfies readonly FounderTaxonomyCategory[];

export const FOUNDER_TAXONOMY_BY_CATEGORY: Readonly<Record<string, FounderTaxonomyCategory>> =
  Object.freeze(Object.fromEntries(
    FOUNDER_TAXONOMY_CONTRACT.map((entry) => [entry.category, entry]),
  ));

export const FOUNDER_QUICK_REVIEW_LABELS_BY_CATEGORY: Readonly<Record<string, readonly string[]>> =
  Object.freeze(Object.fromEntries(
    FOUNDER_TAXONOMY_CONTRACT.map((entry) => [entry.category, entry.quickReviewLabels]),
  ));
'''


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    contract = build_contract(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(render_typescript(args.workbook, contract), encoding="utf-8")
    print(
        f"Generated {args.output} from {args.workbook.name}: "
        f"{len(contract)} categories, "
        f"{sum(len(entry['subcategories']) for entry in contract)} subcategories, "
        f"{sum(len(entry['vibeLabels']) for entry in contract)} VIBES, "
        f"{sum(len(entry['quickReviewLabels']) for entry in contract)} quick reviews."
    )


if __name__ == "__main__":
    main()
